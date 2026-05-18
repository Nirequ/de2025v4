#!/usr/bin/env python3
"""Сборка шаблона ТестКейс.docx со столбцами «Действие», «Ожидаемый
результат», «Результат» (с закладками Результат1..N).

Запуск::

    python module4/docs/_build/build_testcase_docx.py

Результат: ``module4/docs/ТестКейс.docx``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

OUT_PATH = Path(__file__).resolve().parents[1] / "ТестКейс.docx"

TEST_CASES = [
    (
        "Запросить данные через «Получить данные» и отправить корректную "
        "запись (серия 4509, номер 638172).",
        "Поле «Паспорт» — «4509 638172»; результат «Корректный паспорт» "
        "(зелёный); в логе появилась запись IsValid=True.",
    ),
    (
        "Получить запись с буквами в номере (серия 0691, номер 084DH) и "
        "отправить результат.",
        "Результат «Не корректный серия и номер паспорта»; в логе "
        "отображается причина «Номер должен состоять только из цифр».",
    ),
    (
        "Получить запись с короткой серией (например 123) и отправить "
        "результат.",
        "Результат «Не корректный серия и номер паспорта»; в логе "
        "указано «Серия должна содержать ровно 4 цифры».",
    ),
    (
        "Получить запись с числами «0000 / 000000» и отправить результат.",
        "Результат «Не корректный серия и номер паспорта»; в логе "
        "указано «Серия и номер не могут одновременно состоять из нулей».",
    ),
    (
        "Нажать «Отправить результат теста» при выключенном сервере и "
        "оценить поведение приложения.",
        "Приложение не падает, выводит сообщение «Результат не доставлен "
        "на сервер: ...», кнопка «Получить данные» остаётся активной.",
    ),
]

HEADERS = ("№", "Действие", "Ожидаемый результат", "Результат")


def _add_bookmark(paragraph, name: str, bm_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bm_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bm_id))
    paragraph._p.append(start)
    paragraph._p.append(end)


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    run = title.add_run("Тест-кейс: валидация паспортных данных (Модуль 4)")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(
        "Документ оформляется в рамках демо-экзамена 2025, вариант 4. "
        "Поле «Результат» содержит закладки «Результат1..Результат{N}», "
        "которые автоматически заполняет WPF-приложение Module4.App при "
        "нажатии «Отправить результат теста»."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for cell, header in zip(hdr_cells, HEADERS):
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(header)
        run.bold = True
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    widths = (Cm(1.2), Cm(6.0), Cm(6.0), Cm(5.0))
    for column, width in zip(table.columns, widths):
        for cell in column.cells:
            cell.width = width

    for idx, (action, expected) in enumerate(TEST_CASES, start=1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = action
        row[2].text = expected

        result_cell = row[3]
        result_cell.text = ""
        para = result_cell.paragraphs[0]
        _add_bookmark(para, f"Результат{idx}", idx)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note_run = note.add_run(
        "Критерии валидации:\n"
        " 1) серия — ровно 4 цифры, номер — ровно 6 цифр (только цифры);\n"
        " 2) серия и номер не могут одновременно состоять из нулей."
    )
    note_run.italic = True

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"OK -> {path}")
